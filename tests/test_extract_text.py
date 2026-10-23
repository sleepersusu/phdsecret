import sys
from pathlib import Path

from docx import Document
from fpdf import FPDF
from pypdf import PdfReader, PdfWriter

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import extract_text  # noqa: E402

BODY = "Deep learning for early sepsis prediction in intensive care units. " * 3


def make_pdf(path, pages):
    pdf = FPDF()
    pdf.set_font("helvetica", size=11)
    for text in pages:
        pdf.add_page()
        pdf.multi_cell(0, 6, text)
    pdf.output(str(path))
    return path


def lock(src, dst):
    writer = PdfWriter()
    for page in PdfReader(src).pages:
        writer.add_page(page)
    writer.encrypt(user_password="secret", owner_password="owner")
    with open(dst, "wb") as f:
        writer.write(f)


def test_writes_page_marked_text_per_paper(tmp_path):
    (tmp_path / "pdfs").mkdir()
    make_pdf(tmp_path / "pdfs" / "smith2023.pdf", ["First page. " + BODY, "Second page. " + BODY])

    result = extract_text.main([str(tmp_path)])

    text = (tmp_path / "text" / "smith2023.txt").read_text(encoding="utf-8")
    assert result == 0
    assert text.startswith("=== p.1 ===\nFirst page.")
    assert "=== p.2 ===\nSecond page." in text


def test_skips_password_protected_paper(tmp_path, capsys):
    (tmp_path / "pdfs").mkdir()
    plain = make_pdf(tmp_path / "plain.pdf", [BODY])
    lock(plain, tmp_path / "pdfs" / "locked.pdf")

    extract_text.main([str(tmp_path)])

    assert not (tmp_path / "text" / "locked.txt").exists()
    assert "locked.pdf" in capsys.readouterr().err


def test_copies_markdown_proposal(tmp_path):
    (tmp_path / "pdfs").mkdir()
    proposal = tmp_path / "proposal.md"
    proposal.write_text("# Aims\nWe will predict sepsis.\n", encoding="utf-8")

    extract_text.main([str(tmp_path), "--proposal", str(proposal)])

    assert (tmp_path / "text" / "proposal.txt").read_text(encoding="utf-8") == "# Aims\nWe will predict sepsis.\n"


def test_docx_proposal_keeps_headings(tmp_path):
    (tmp_path / "pdfs").mkdir()
    doc = Document()
    doc.add_heading("Background", level=1)
    doc.add_paragraph("Sepsis kills millions every year.")
    doc.add_heading("Hypothesis", level=2)
    doc.save(tmp_path / "proposal.docx")

    extract_text.main([str(tmp_path), "--proposal", str(tmp_path / "proposal.docx")])

    text = (tmp_path / "text" / "proposal.txt").read_text(encoding="utf-8")
    assert text == "# Background\nSepsis kills millions every year.\n## Hypothesis\n"


def test_pdf_proposal_drops_page_markers(tmp_path):
    (tmp_path / "pdfs").mkdir()
    make_pdf(tmp_path / "proposal.pdf", ["Aims. " + BODY, "Methods. " + BODY])

    extract_text.main([str(tmp_path), "--proposal", str(tmp_path / "proposal.pdf")])

    text = (tmp_path / "text" / "proposal.txt").read_text(encoding="utf-8")
    assert "=== p." not in text
    assert "Aims." in text and "Methods." in text
